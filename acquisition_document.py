"""Compact, editable commercial acquisition report export.

The document uses the same paid brief as HTML. Source extracts remain in the
online evidence appendix. Material risks are retained even when a complex
property legitimately needs more than the standard five-page structure.
"""
from io import BytesIO
import re

from acquisition_report import commercial_brief


def render_docx(report):
    """Return DOCX bytes for a full report; a free snapshot cannot be exported."""
    if report.get('access') == 'snapshot':
        raise ValueError('A full report is required for the commercial document')
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    brief = commercial_brief(report)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.65)
    section.left_margin = section.right_margin = Cm(1.7)
    section.header_distance = section.footer_distance = Cm(.7)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2',
                 'Heading 3', 'Header', 'Footer'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        # Remove template theme colours so Word/LibreOffice agree.
        for color in style.element.xpath('.//w:color'):
            for attr in ('themeColor', 'themeTint', 'themeShade'):
                color.attrib.pop(qn('w:' + attr), None)
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    normal = doc.styles['Normal']
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True
    for name, size, before, after in [('Title', 23, 0, 10),
                                     ('Heading 1', 19, 0, 13),
                                     ('Heading 2', 12, 11, 6),
                                     ('Heading 3', 10.5, 6, 4)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    def clean(value):
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value if value is not None else ''))

    def paragraph(text, style=None, size=None, bold=False, after=None):
        p = doc.add_paragraph(style=style)
        run = p.add_run(clean(text))
        if size:
            run.font.size = Pt(size)
        if bold:
            run.bold = True
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    def note(text):
        if text:
            p = paragraph(text, size=8, after=4)
            p.paragraph_format.line_spacing = 1.08
            p.runs[0].font.color.rgb = RGBColor.from_string('52616B')
            return p

    def heading(text, level=1):
        return doc.add_heading(clean(text), level=level)

    def shade(cell, fill):
        element = OxmlElement('w:shd')
        element.set(qn('w:fill'), fill)
        cell._tc.get_or_add_tcPr().append(element)

    def table(headers, rows, widths, numeric=()):
        if not rows:
            return None
        table = doc.add_table(rows=1, cols=len(headers))
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement('w:' + edge)
            for name, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                border.set(qn('w:' + name), value)
            borders.append(border)
        tbl_pr.append(borders)
        margins = OxmlElement('w:tblCellMar')
        for edge, value in [('top', '85'), ('bottom', '85'), ('left', '100'), ('right', '100')]:
            item = OxmlElement('w:' + edge)
            item.set(qn('w:w'), value)
            item.set(qn('w:type'), 'dxa')
            margins.append(item)
        tbl_pr.append(margins)
        repeat = OxmlElement('w:tblHeader')
        repeat.set(qn('w:val'), 'true')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        all_rows = [headers] + rows
        for ri, values in enumerate(all_rows):
            row = table.rows[0] if ri == 0 else table.add_row()
            no_split = OxmlElement('w:cantSplit')
            row._tr.get_or_add_trPr().append(no_split)
            for ci, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                if ci in numeric:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                run = p.add_run(clean(value))
                run.font.size = Pt(8.5 if len(headers) >= 4 else 9)
                if ri == 0:
                    shade(cell, '203747')
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                elif ri % 2 == 0:
                    shade(cell, 'F3F5F7')
        paragraph('', size=3, after=2)
        return table

    # A discreet running footer helps a printed report remain identifiable.
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.space_before = Pt(3)
    footer.paragraph_format.tab_stops.add_tab_stop(Cm(17.6), WD_ALIGN_PARAGRAPH.RIGHT)
    run = footer.add_run('Auction Sniper Acquisition Intelligence')
    run.font.size = Pt(7.5)
    footer.add_run('\t')
    for index, instruction in enumerate(('PAGE', 'NUMPAGES')):
        if index:
            footer.add_run(' / ').font.size = Pt(7.5)
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), instruction)
        footer._p.append(field)
    doc.core_properties.title = clean(report.get('property') or 'Property acquisition report')
    doc.core_properties.subject = 'Commercial property acquisition review'
    doc.core_properties.author = 'Auction Sniper'
    doc.core_properties.keywords = 'acquisition, commercial property, legal pack'

    paragraph('AUCTION SNIPER  ACQUISITION INTELLIGENCE', size=8, bold=True, after=8)
    paragraph(report.get('property') or 'Property acquisition report', style='Title')
    note(str(brief.get('as_of') or '') + '  |  ' + str(brief.get('status') or 'Acquisition review'))
    paragraph(brief.get('summary'), size=10.5)
    heading('Acquisition conclusion', 2)
    paragraph(brief.get('conclusion'))
    figures = [[r.get('label'), r.get('value'), r.get('basis')] for r in brief.get('key_figures', [])]
    table(['Key figure', 'Value', 'Basis'], figures, [4.2, 4.4, 9.0])
    risks = brief.get('risks', [])
    if risks:
        heading('Decision priorities', 2)
        for i, risk in enumerate(risks[:3], 1):
            p = paragraph('', after=5)
            p.add_run(str(i) + ' ' + clean(risk.get('title')) + '. ').bold = True
            p.add_run(clean(risk.get('action')))
    note(brief.get('scope'))

    heading('Income and lease terms').paragraph_format.page_break_before = True
    income_rows = []
    for r in brief.get('income_rows', []):
        unit = clean(r.get('unit'))
        if r.get('source'):
            unit += '\n' + clean(r['source'])
        income_rows.append([unit, r.get('tenant'), r.get('rent'), r.get('terms')])
    table(['Accommodation', 'Legal tenant', 'Annual rent', 'Terms and qualifications'],
          income_rows, [3.6, 3.7, 2.3, 8.0], numeric=(2,))
    for text in brief.get('income_commentary', []):
        if text:
            paragraph(text)

    heading('Price and acquisition costs').paragraph_format.page_break_before = True
    table(['Cost or cash item', 'Amount', 'Basis and treatment'],
          [[r.get('item'), r.get('amount'), r.get('basis')] for r in brief.get('cost_rows', [])],
          [5.1, 3.0, 9.5], numeric=(1,))
    for text in brief.get('financial_commentary', []):
        if text:
            paragraph(text)
    if brief.get('sensitivity_rows'):
        heading('Income sensitivity', 2)
        table(['Scenario', 'Annual rent', 'Gross yield', 'Interpretation'],
              [[r.get('scenario'), r.get('rent'), r.get('yield'), r.get('meaning')]
               for r in brief['sensitivity_rows']], [4.1, 2.8, 2.6, 8.1], numeric=(1, 2))

    heading('Prioritised acquisition risks').paragraph_format.page_break_before = True
    if not risks:
        paragraph('No specific risks were established by this review. Absence of an extracted finding is not a clean legal or condition opinion.')
    for i, risk in enumerate(risks, 1):
        heading(str(i) + ' ' + clean(risk.get('title')), 3)
        p = paragraph('', after=5)
        if risk.get('priority'):
            p.add_run(clean(risk['priority']) + '. ').bold = True
        p.add_run(clean(risk.get('finding')))
        p.paragraph_format.keep_with_next = True
        p = paragraph('', after=3)
        p.add_run('Action. ').bold = True
        p.add_run(clean(risk.get('action')))
        p.paragraph_format.keep_with_next = True
        note(risk.get('source'))

    heading('Before an unconditional commitment').paragraph_format.page_break_before = True
    for check in brief.get('checks', []):
        heading(check.get('who') or 'Required check', 3)
        paragraph(check.get('action'), after=4)
        note(check.get('source'))
    heading('Evidence and scope', 2)
    for text in brief.get('source_notes', []):
        note(text)
    heading('Acquisition conclusion', 2)
    paragraph(brief.get('conclusion'))
    note(brief.get('scope') or report.get('disclaimer'))
    engine = report.get('engine') or {}
    note('Report ' + str(report.get('report_id', '')) +
         (' | Engine ' + str(engine.get('release', '')) + ' / ' + str(engine.get('build', '')) if engine else ''))
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
