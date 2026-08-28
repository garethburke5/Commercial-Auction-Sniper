"""Mixed-format, content-led ingestion for Auction Sniper Buyer Due Diligence.

A legal-pack upload is treated as an arbitrary container of evidence, not a fixed
set of filenames or extensions. Files are sniffed from their bytes, ZIPs are
expanded safely, machine-readable formats are extracted, and visual/binary
items are inventoried rather than silently discarded. No OCR is performed here.
"""
from __future__ import annotations
from dataclasses import dataclass,field
from io import BytesIO
from pathlib import Path
from typing import Iterable
from email import policy
from email.parser import BytesParser
from html import unescape
from zipfile import ZipFile,BadZipFile
import csv,hashlib,json,re
from legal_pack_engine import PackDocument,make_document
try:
 from pypdf import PdfReader
except Exception:PdfReader=None
try:
 from docx import Document
except Exception:Document=None
try:
 from openpyxl import load_workbook
except Exception:load_workbook=None

MAX_ARCHIVE_MEMBERS=250
MAX_ARCHIVE_UNCOMPRESSED=100*1024*1024
TEXT_EXTS={".txt",".csv",".tsv",".json",".xml",".html",".htm",".eml",".rtf"}
IMAGE_EXTS={".png",".jpg",".jpeg",".webp",".tif",".tiff",".bmp",".gif"}

@dataclass
class IngestIssue: filename:str;code:str;message:str;severity:str="warning"
@dataclass
class IngestAsset:
 filename:str;detected_type:str;status:str;message:str="";sha256:str="";bytes:int=0;container:str|None=None
@dataclass
class IngestResult:
 documents:list[PackDocument]=field(default_factory=list);issues:list[IngestIssue]=field(default_factory=list);duplicates:list[str]=field(default_factory=list);assets:list[IngestAsset]=field(default_factory=list);total_bytes:int=0;total_pages:int=0;expanded_files:int=0

def _clean(text):return "\n".join(line.rstrip() for line in (text or "").replace("\x00","").splitlines()).strip()
def _sha(raw):return hashlib.sha256(raw).hexdigest()
def _looks_text(raw):
 if not raw:return True
 sample=raw[:4096]
 if b"\x00" in sample:return False
 printable=sum(1 for b in sample if b in b"\t\n\r" or 32<=b<=126 or b>=128)
 return printable/max(1,len(sample))>.90

def detect_type(filename,raw):
 """Sniff actual content first; extension is only a secondary hint."""
 ext=Path(filename).suffix.lower();head=raw[:16]
 if head.startswith(b"%PDF-"):return "pdf"
 if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):return "legacy_office"
 if head.startswith(b"\x89PNG\r\n\x1a\n"):return "image"
 if head[:3]==b"\xff\xd8\xff" or head.startswith((b"GIF87a",b"GIF89a",b"BM",b"II*\x00",b"MM\x00*",b"RIFF")):return "image"
 if head.startswith(b"PK\x03\x04"):
  try:
   with ZipFile(BytesIO(raw)) as z:
    names=set(z.namelist())
    if "word/document.xml" in names:return "docx"
    if "xl/workbook.xml" in names:return "xlsx"
    return "zip"
  except BadZipFile:return "binary"
 if ext in TEXT_EXTS or _looks_text(raw):return "text"
 if ext in IMAGE_EXTS:return "image"
 return "binary"

def extract_pdf(raw):
 if PdfReader is None:raise RuntimeError("pypdf is not installed")
 reader=PdfReader(BytesIO(raw));pages=[];chunks=[]
 for i,page in enumerate(reader.pages,1):
  txt=_clean(page.extract_text() or "");pages.append({"page":i,"text":txt,"chars":len(txt)})
  if txt:chunks.append(f"\n--- PAGE {i} ---\n{txt}")
 return "".join(chunks).strip(),{"pages":pages,"page_count":len(reader.pages),"extraction":"pypdf"}
def extract_docx(raw):
 if Document is None:raise RuntimeError("python-docx is not installed")
 doc=Document(BytesIO(raw));blocks=[];paragraphs=[]
 for i,p in enumerate(doc.paragraphs,1):
  txt=_clean(p.text)
  if txt:paragraphs.append({"paragraph":i,"text":txt});blocks.append(f"--- PARAGRAPH {i} ---\n{txt}")
 tables=[]
 for ti,table in enumerate(doc.tables,1):
  rows=[]
  for ri,row in enumerate(table.rows,1):
   vals=[_clean(c.text) for c in row.cells];line=" | ".join(vals);rows.append({"row":ri,"cells":vals})
   if line.strip(" |"):blocks.append(f"--- TABLE {ti} ROW {ri} ---\n{line}")
  tables.append({"table":ti,"rows":rows})
 return "\n".join(blocks),{"paragraphs":paragraphs,"tables":tables,"extraction":"python-docx"}
def extract_xlsx(raw):
 if load_workbook is None:raise RuntimeError("openpyxl is not installed")
 wb=load_workbook(BytesIO(raw),read_only=True,data_only=True);blocks=[];sheets=[]
 for ws in wb.worksheets:
  rows=[]
  for ri,row in enumerate(ws.iter_rows(values_only=True),1):
   vals=["" if v is None else str(v) for v in row]
   if any(v.strip() for v in vals):
    rows.append({"row":ri,"cells":vals});blocks.append(f"--- SHEET {ws.title} ROW {ri} ---\n"+" | ".join(vals))
  sheets.append({"sheet":ws.title,"rows":rows})
 return "\n".join(blocks),{"sheets":sheets,"extraction":"openpyxl"}
def extract_text(filename,raw):
 ext=Path(filename).suffix.lower()
 for enc in ("utf-8","utf-8-sig","cp1252","latin-1"):
  try:text=raw.decode(enc);break
  except UnicodeDecodeError:continue
 else:text=raw.decode("utf-8","replace");enc="replacement"
 if ext==".html" or ext==".htm":text=re.sub(r"<[^>]+>"," ",unescape(text))
 elif ext==".eml":
  msg=BytesParser(policy=policy.default).parsebytes(raw);parts=[]
  for p in msg.walk():
   if p.get_content_type()=="text/plain":
    try:parts.append(p.get_content())
    except Exception:pass
  text="\n".join(parts) or text
 elif ext==".rtf":text=re.sub(r"\\[a-z]+\d* ?|[{}]"," ",text)
 return _clean(text),{"extraction":f"text/{enc}"}
def locate_evidence(metadata,needle):
 if not needle:return {}
 target=re.sub(r"\s+"," ",needle).strip().lower()
 def hit(text):return target in re.sub(r"\s+"," ",text or "").lower()
 for p in metadata.get("pages",[]):
  if hit(p.get("text","")):return {"page":p["page"]}
 for p in metadata.get("paragraphs",[]):
  if hit(p.get("text","")):return {"paragraph":p["paragraph"]}
 for t in metadata.get("tables",[]):
  for row in t.get("rows",[]):
   if hit(" | ".join(row.get("cells",[]))):return {"table":t["table"],"row":row["row"]}
 for s in metadata.get("sheets",[]):
  for row in s.get("rows",[]):
   if hit(" | ".join(row.get("cells",[]))):return {"sheet":s["sheet"],"row":row["row"]}
 return {}
def ingest_bytes(filename,raw):
 kind=detect_type(filename,raw)
 if kind=="legacy_office":raise ValueError("Legacy Office/OLE file detected. Convert .doc/.xls to PDF, DOCX or XLSX, or route through the conversion service before analysis.")
 if kind in ("zip","image","binary"):raise ValueError(f"{kind.upper()} requires pack-level handling rather than direct text ingestion.")
 if kind=="pdf":text,meta=extract_pdf(raw)
 elif kind=="docx":text,meta=extract_docx(raw)
 elif kind=="xlsx":text,meta=extract_xlsx(raw)
 else:text,meta=extract_text(filename,raw)
 meta.update({"filename":filename,"detected_type":kind,"extension":Path(filename).suffix.lower(),"bytes":len(raw),"sha256":_sha(raw),"text_chars":len(text),"text_empty":not bool(text.strip())})
 return make_document(filename,text,raw=raw,metadata=meta)

def ingest_pack(files:Iterable[tuple[str,bytes]]):
 result=IngestResult();seen=set()
 def process(filename,raw,container=None,depth=0):
  digest=_sha(raw);kind=detect_type(filename,raw)
  if digest in seen:result.duplicates.append(filename);return
  seen.add(digest);result.expanded_files+=1
  if kind=="zip":
   if depth>=3:
    result.issues.append(IngestIssue(filename,"ARCHIVE_DEPTH","Nested archive depth limit reached."));return
   try:
    with ZipFile(BytesIO(raw)) as z:
     members=[i for i in z.infolist() if not i.is_dir()]
     if len(members)>MAX_ARCHIVE_MEMBERS or sum(i.file_size for i in members)>MAX_ARCHIVE_UNCOMPRESSED:
      result.issues.append(IngestIssue(filename,"ARCHIVE_LIMIT","Archive exceeds safe member/size limits.","error"));return
     result.assets.append(IngestAsset(filename,"zip","expanded",f"Expanded {len(members)} member(s).",digest,len(raw),container))
     for info in members:
      member=f"{filename}::{info.filename}";process(member,z.read(info),filename,depth+1)
   except Exception as e:result.issues.append(IngestIssue(filename,"ARCHIVE_FAILED",f"Archive expansion failed: {e}","error"))
   return
  if kind=="image":
   result.assets.append(IngestAsset(filename,"image","visual_review_required","Image/scan retained for visual analysis; no OCR was attempted.",digest,len(raw),container));return
  if kind=="binary":
   result.assets.append(IngestAsset(filename,"binary","unsupported","Binary item retained in inventory but not machine-read.",digest,len(raw),container));result.issues.append(IngestIssue(filename,"UNSUPPORTED_BINARY","File could not be safely interpreted as a supported machine-readable format."));return
  if kind=="legacy_office":
   result.assets.append(IngestAsset(filename,"legacy_office","conversion_required","Legacy Office file requires conversion before text analysis.",digest,len(raw),container));result.issues.append(IngestIssue(filename,"LEGACY_OFFICE","Legacy .doc/.xls/OLE file requires conversion before analysis."));return
  try:
   d=ingest_bytes(filename,raw);d.metadata["container"]=container;result.documents.append(d);result.total_pages+=int(d.metadata.get("page_count",0) or 0)
   result.assets.append(IngestAsset(filename,kind,"analysed" if d.text.strip() else "partial",sha256=digest,bytes=len(raw),container=container))
   if d.metadata.get("text_empty"):result.issues.append(IngestIssue(filename,"NO_EXTRACTABLE_TEXT","No machine-readable text was extracted; retain for visual review."))
  except Exception as e:result.issues.append(IngestIssue(filename,"EXTRACTION_FAILED",f"Document extraction failed: {e}","error"))
 for filename,raw in files:
  result.total_bytes+=len(raw);process(filename,raw)
 return result
