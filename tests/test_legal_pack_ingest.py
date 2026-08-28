from io import BytesIO
from zipfile import ZipFile
from docx import Document
from openpyxl import Workbook
from legal_pack_ingest import ingest_bytes,ingest_pack,detect_type
from legal_pack_engine import DocType

def make_docx_bytes():
 bio=BytesIO();d=Document();d.add_paragraph('Special Conditions of Sale');d.add_paragraph('rent of £20,000 per annum and a term of five years');table=d.add_table(rows=1,cols=2);table.cell(0,0).text='Tenant';table.cell(0,1).text='Boots Opticians Professional Services Ltd';d.save(bio);return bio.getvalue()
def make_xlsx_bytes():
 bio=BytesIO();w=Workbook();s=w.active;s.title='Arrears';s.append(['Tenant','Balance']);s.append(['Boots',0]);w.save(bio);return bio.getvalue()
def test_docx_ingestion_preserves_paragraphs_and_tables():
 raw=make_docx_bytes();doc=ingest_bytes('Special conditions - 8 Red Street.docx',raw)
 assert doc.doc_type==DocType.SPECIAL_CONDITIONS and '£20,000' in doc.text and 'Boots Opticians' in doc.text
 assert doc.metadata['tables'][0]['rows'][0]['cells'][0]=='Tenant' and doc.metadata['sha256']==doc.sha256
def test_content_sniffing_beats_wrong_extension():
 raw=make_docx_bytes();doc=ingest_bytes('mystery.bin',raw)
 assert doc.doc_type==DocType.SPECIAL_CONDITIONS and doc.metadata['detected_type']=='docx'
def test_mixed_zip_expands_and_inventories_visuals_and_spreadsheet():
 bio=BytesIO()
 with ZipFile(bio,'w') as z:
  z.writestr('lease/Special Conditions.docx',make_docx_bytes());z.writestr('accounts/arrears.xlsx',make_xlsx_bytes());z.writestr('plans/site-plan.png',b'\x89PNG\r\n\x1a\n'+b'x'*20)
 result=ingest_pack([('pack.zip',bio.getvalue())])
 assert len(result.documents)==2
 assert any(d.metadata.get('detected_type')=='xlsx' for d in result.documents)
 assert any(a.status=='visual_review_required' for a in result.assets)
 assert any(a.status=='expanded' for a in result.assets)
def test_pack_deduplicates_and_flags_real_legacy_office_signature():
 raw=make_docx_bytes();legacy=b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'+b'legacy'
 result=ingest_pack([('a.docx',raw),('copy.docx',raw),('CPSE2.doc',legacy)])
 assert len(result.documents)==1 and result.duplicates==['copy.docx']
 assert any(i.code=='LEGACY_OFFICE' for i in result.issues)
 assert any(a.status=='conversion_required' for a in result.assets)
def test_text_document_classifies():
 doc=ingest_bytes('EPC.txt',b'Energy rating C (58)');assert doc.doc_type==DocType.EPC
